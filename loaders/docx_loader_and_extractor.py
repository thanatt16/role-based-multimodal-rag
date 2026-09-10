from docx import Document
import os
import base64
import zipfile

from openai import OpenAI
from dotenv import load_dotenv


load_dotenv()

client = OpenAI()

def extract_image_content(image_bytes, image_extension):

    base64_image = base64.b64encode(image_bytes).decode("utf-8")

    if image_extension.lower() in ["jpg", "jpeg"]:
        mime_type = "image/jpeg"
    else:
        mime_type = f"image/{image_extension.lower()}"

    response = client.responses.create(
        model="gpt-5.6-luna",
        input=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "input_text",
                        "text": """
Extract all useful information from this image for a RAG knowledge base.

Read all visible text accurately.
Preserve important numbers, names, labels and values exactly as shown.

If there is a table:
- Extract its headers, rows and values.

If there is a chart:
- Extract visible labels, axes and values.
- Describe clearly visible trends and relationships.
- Do not invent precise values that are not explicitly visible.

If there is a diagram:
- Describe its components and relationships.

Do not invent information that is not visible.
"""
                    },
                    {
                        "type": "input_image",
                        "image_url": f"data:{mime_type};base64,{base64_image}",
                        "detail": "high"
                    }
                ]
            }
        ]
    )

    return response.output_text


def table_to_text(table):

    rows = []

    for row in table.rows:

        values = []

        for cell in row.cells:
            values.append(cell.text.strip())

        row_text = " | ".join(values)
        rows.append(row_text)

    return "\n".join(rows)

def load_docx(path):

    doc = Document(path)

    documents = []

    paragraphs = []

    for paragraph in doc.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    full_text = "\n".join(paragraphs)

    text_document = {
        "content": full_text,
        "metadata": {
            "source": os.path.basename(path),
            "file_type": "docx",
            "content_type": "text"
        }
    }

    documents.append(text_document)

    for table_index, table in enumerate(doc.tables, start=1):

        table_text = table_to_text(table)

        table_document = {
            "content": table_text,
            "metadata": {
                "source": os.path.basename(path),
                "file_type": "docx",
                "content_type": "table",
                "table_index": table_index
            }
        }

        documents.append(table_document)


    with zipfile.ZipFile(path, "r") as docx_zip:

        image_files = [
            file_name
            for file_name in docx_zip.namelist()
            if file_name.startswith("word/media/")
        ]

        for image_index, image_file in enumerate(image_files, start=1):

            image_bytes = docx_zip.read(image_file)

            image_extension = os.path.splitext(image_file)[1].replace(".", "")

            vlm_result = extract_image_content(image_bytes, image_extension)

            image_document = {
                "content": vlm_result,
                "metadata": {
                    "source": os.path.basename(path),
                    "file_type": "docx",
                    "content_type": "image",
                    "image_index": image_index,
                    "image_name": os.path.basename(image_file)
                }
            }

            documents.append(image_document)

    return documents

def load_docx_folder(folder_path):

    all_documents = []

    for filename in os.listdir(folder_path):

        if filename.lower().endswith(".docx"):

            file_path = os.path.join(folder_path, filename)

            documents = load_docx(file_path)

            all_documents.extend(documents)

    return all_documents

if __name__ == "__main__":

    docx_folder = r"E:\Machine_learning\AI\RAG\PROJECT\data\docx_data"

    documents = load_docx_folder(docx_folder)

    print("Number of DOCX documents:", len(documents))

    for document in documents:

        print(document["metadata"])
        print(document["content"][:300])
        print("----------------------")