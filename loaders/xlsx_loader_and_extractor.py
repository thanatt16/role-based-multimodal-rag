from openpyxl import load_workbook
import os
import base64

from openai import OpenAI
from dotenv import load_dotenv


load_dotenv()

client = OpenAI()


def extract_image_content(image_bytes, image_extension):

    base64_image = base64.b64encode(image_bytes).decode("utf-8")

    if image_extension.lower() in ["jpg", "jpeg"]:
        mime_type = "image/jpeg"
    else:
        mime_type = f"image/{image_extension.lower()}"

    response = client.responses.create(
        model="gpt-5.6-luna",
        input=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "input_text",
                        "text": """
Extract all useful information from this image for a RAG knowledge base.

Read all visible text accurately.
Preserve important numbers, names, labels and values exactly as shown.

If there is a table:
- Extract its headers, rows and values.

If there is a chart:
- Extract visible labels, axes and values.
- Describe clearly visible trends and relationships.
- Do not invent precise values that are not explicitly visible.

If there is a diagram:
- Describe its components and relationships.

Do not invent information that is not visible.
"""
                    },
                    {
                        "type": "input_image",
                        "image_url": f"data:{mime_type};base64,{base64_image}",
                        "detail": "high"
                    }
                ]
            }
        ]
    )

    return response.output_text

def load_xlsx(path):

    workbook = load_workbook(path, data_only=True)

    documents = []

    for sheet in workbook.worksheets:

        rows = []

        for row in sheet.iter_rows(values_only=True):

            values = []

            for value in row:

                if value is not None:
                    values.append(str(value))
                else:
                    values.append("")

            row_text = " | ".join(values)
            rows.append(row_text)

        sheet_text = "\n".join(rows)

        table_document = {
            "content": sheet_text,
            "metadata": {
                "source": os.path.basename(path),
                "file_type": "xlsx",
                "content_type": "table",
                "sheet": sheet.title
            }
        }

        documents.append(table_document)


        for image_index, image in enumerate(sheet._images, start=1):

            image_bytes = image._data()
            image_extension = image.format

            vlm_result = extract_image_content(image_bytes, image_extension)

            image_document = {
                "content": vlm_result,
                "metadata": {
                    "source": os.path.basename(path),
                    "file_type": "xlsx",
                    "content_type": "image",
                    "sheet": sheet.title,
                    "image_index": image_index,
                    "width": image.width,
                    "height": image.height
                }
            }

            documents.append(image_document)

    return documents



def load_xlsx_folder(folder_path):

    all_documents = []

    for filename in os.listdir(folder_path):

        if filename.lower().endswith(".xlsx"):

            file_path = os.path.join(folder_path, filename)

            documents = load_xlsx(file_path)

            all_documents.extend(documents)

    return all_documents


if __name__ == "__main__":

    xlsx_folder = "data/xlsx_data"

    documents = load_xlsx_folder(xlsx_folder)

    print("Number of XLSX documents:", len(documents))

    for document in documents:

        print(document["metadata"])
        print(document["content"][:300])
        print("----------------------")